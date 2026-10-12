"""usage: python make_projection.py MODEL [input.xlsx]
Reads analysis_out/count_index_all_MODEL.xlsx (Sum rows) -> analysis_out/projection_MODEL.xlsx (difference-method projection, Task 1 + Task 2)"""
import sys, os, re, numpy as np, openpyxl
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImg
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

model = sys.argv[1]; src = sys.argv[2] if len(sys.argv) > 2 else f'analysis_out/count_index_all_{model}.xlsx'
OUT = f'analysis_out/projection_{model}.xlsx'; PD = f'analysis_out/proj_plots_{model}'; os.makedirs(PD, exist_ok=True)
KS = np.array([10, 20, 30, 40, 50.]); GRID = np.arange(60, 1001, 10); REP = [100, 200, 300, 400, 500, 1000]

# ---- 1. read the Sum row of every {dataset}_K{K} sheet:  S[ds][K] = {size: count}, NQ[ds]
S, NQ = {}, {}
wbi = openpyxl.load_workbook(src, read_only=True, data_only=True)
for sh in wbi.sheetnames:
    mm = re.match(r'^(.+)_K(\d+)$', sh)
    if not mm: continue
    ds, K = mm.group(1), int(mm.group(2))
    it = wbi[sh].iter_rows(values_only=True); hdr = [str(h) for h in next(it)]; nq = 0; last = None
    for r in it:
        if r[0] is None: continue
        if r[0] == 'Sum': last = r
        else: nq += 1
    S.setdefault(ds, {})[K] = {s: int(last[hdr.index(str(s))]) for s in range(1, K + 1)}; NQ[ds] = nq

# ---- 2. difference-method fit
def fit(K, D):
    m = D > 0
    if m.sum() < 2: return np.nan, np.nan, int(m.sum()), np.nan
    b, la = np.polyfit(np.log(K[m]), np.log(D[m]), 1); y = np.log(D[m]); f = la + b * np.log(K[m])
    return np.exp(la), b, int(m.sum()), 1 - ((y - f) ** 2).sum() / max(((y - y.mean()) ** 2).sum(), 1e-12)
def project(y50, A, b, sg):
    out, y = {}, y50
    for K in GRID: y = max(y + sg * A * K ** b, 0.0); out[K] = y
    return out
def limit(y50, A, b, sg): return np.nan if not b < -1 else max(y50 + sg * (A / 10) * 55 ** (b + 1) / (-(b + 1)), 0.0)
def status(A, b):
    if A != A: return 'no fit (<2 usable differences)'
    if b < -1: return 'converges (increments shrink fast)'
    if b < 0: return 'slow: increments shrink, total keeps changing'
    return 'unreliable (fitted increments do not shrink)'
def analyse(y, sg):
    D = np.diff(y) * sg; A, b, n, r2 = fit(KS[1:], D); A3, b3, n3, _ = fit(KS[1:4], D[:3])
    ho = abs(y[3] + sg * A3 * 50 ** b3 - y[4]) / max(y[4], 1) if n3 >= 2 else np.nan
    ok = A == A
    return dict(y=y, D=np.diff(y), A=A, b=b, n=n, r2=r2, ho=ho, pr=project(y[4], A, b, sg) if ok else {k: np.nan for k in GRID},
                lim=limit(y[4], A, b, sg) if ok else np.nan, st=status(A, b))
def curve(ax, r, title, ylab):
    ax.plot(KS, r['y'], 'ko', ms=4)
    if r['A'] == r['A']:
        ax.plot([50] + list(GRID), [r['y'][4]] + [r['pr'][k] for k in GRID], 'r-', lw=1.2); ax.plot(REP, [r['pr'][k] for k in REP], 'rs', ms=3)
        if r['lim'] == r['lim']: ax.axhline(r['lim'], c='gray', ls='--', lw=.8)
    ax.set_xscale('log'); ax.set_title(title, fontsize=9); ax.tick_params(labelsize=7); ax.set_xlabel('K', fontsize=7); ax.set_ylabel(ylab, fontsize=7)

# ---- 3. workbook
HDR = PatternFill('solid', fgColor='DDE6F0'); B = Font(bold=True)
cols = ['y(K=10)', 'y(K=20)', 'y(K=30)', 'y(K=40)', 'y(K=50)', 'D20', 'D30', 'D40', 'D50', 'A', 'b', 'R2 (log-log)', 'usable diffs',
        'hold-out err K50', 'limit (if converges)', 'status'] + [f'K={k}' for k in REP]
vals = lambda lab, r: [lab, *r['y'], *r['D'], r['A'], r['b'], r['r2'], r['n'], r['ho'], r['lim'], r['st'], *[r['pr'][k] for k in REP]]
def head(w):
    w.append(['series', *cols]); [setattr(c, 'font', B) or setattr(c, 'fill', HDR) for c in w[w.max_row]]
wb = Workbook(); rd = wb.active; rd.title = 'README'
for i, t in enumerate([f'Projection workbook - {model}', '',
    'y_s(K) = sum over questions of the number of clusters of size s (column s) at K samples. Increment D_s(K) = y_s(K) - y_s(K-10), K=20..50.',
    'Fit D_s(K) = A*K^b (log-log least squares on the positive increments). Projection: y_s(K+10) = y_s(K) + A*(K+10)^b up to K=1000.',
    'Closed form: y_s(N) ~ y_s(50) + (A/10)*(N^(b+1) - 55^(b+1))/(b+1).  b < -1: converges to a limit; -1 <= b < 0: increments shrink but total keeps growing; b >= 0: unusable.',
    'Task 2 (most confident): g(K) = number of questions whose K samples all fall in one cluster (column K at K); drops are fitted the same way and subtracted.',
    'hold-out err K50: relative error predicting y(50) from a fit on K=10..40 only. Only 4 increments per fit; negative increments are dropped, so columns 3-9 are unreliable.']):
    rd.cell(i + 1, 1, t)
rd['A1'].font = Font(bold=True, size=14); rd.column_dimensions['A'].width = 150
summ = wb.create_sheet('Summary'); summ.append(['dataset', *['series', *cols]]); [setattr(c, 'font', B) or setattr(c, 'fill', HDR) for c in summ[1]]
for ds in sorted(S):
    nq = NQ[ds]; w = wb.create_sheet(ds)
    w.append([f'{model} / {ds}  ({nq} questions)  - Task 1: columns 1-9']); w['A1'].font = Font(bold=True, size=12); head(w)
    res = {}
    for s in range(1, 10):
        res[s] = analyse(np.array([S[ds][k][s] for k in KS.astype(int)], float), +1); w.append(vals(f'column {s}', res[s])); summ.append([ds, *vals(f'column {s}', res[s])])
    rt = analyse(np.array([S[ds][k][k] for k in KS.astype(int)], float), -1)
    w.append([]); w.append(['Task 2: most confident (column K at K), count of questions']); w.cell(w.max_row, 1).font = Font(bold=True, size=12); head(w)
    w.append(vals('column K at K', rt)); summ.append([ds, *vals('most confident', rt)])
    w.append(['as fraction of questions', *(rt['y'] / nq), *[None] * 10, rt['lim'] / nq if rt['lim'] == rt['lim'] else None, None, *[rt['pr'][k] / nq for k in REP]])
    for x in (w, summ):
        for row in x.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float): c.number_format = '0.000' if abs(c.value) < 10 else '#,##0.0'
        for i in range(1, 30): x.column_dimensions[L(i)].width = 13
        x.column_dimensions['A'].width = 26; x.column_dimensions['Q'].width = 40
    fig, ax = plt.subplots(3, 3, figsize=(13, 9))
    for s, a in zip(range(1, 10), ax.ravel()): curve(a, res[s], f'column {s}: b={res[s]["b"]:.2f}' if res[s]['A'] == res[s]['A'] else f'column {s}: no fit', 'column sum')
    fig.suptitle(f'{model} / {ds}: columns 1-9 (black = observed K=10..50, red = projection to K=1000, squares = K=100..500,1000)', fontsize=10)
    fig.tight_layout(); p1 = f'{PD}/{ds}_cols.png'; fig.savefig(p1, dpi=75); plt.close(fig)
    fig, a = plt.subplots(figsize=(7, 4.2)); curve(a, rt, f'Task 2 most confident: b={rt["b"]:.2f}', 'questions with all K samples in one cluster')
    fig.tight_layout(); p2 = f'{PD}/{ds}_confident.png'; fig.savefig(p2, dpi=85); plt.close(fig)
    w.add_image(XImg(p1), 'A16'); w.add_image(XImg(p2), 'A56')
wb._sheets.insert(1, wb._sheets.pop(wb._sheets.index(summ)))
wb.save(OUT); print('saved', OUT)